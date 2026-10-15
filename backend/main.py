import io
import csv
from datetime import datetime, timedelta
from typing import Optional, List
from contextlib import asynccontextmanager

from fastapi import FastAPI, Query, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

import backend.database as db
import backend.utils_backup as backup
from backend.config import obter_diretorio_dist

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    TEM_OPENPYXL = True
except ImportError:
    TEM_OPENPYXL = False

@asynccontextmanager
async def lifespan(app: FastAPI):
    db.criar_tabelas()
    backup.fazer_backup_banco()
    yield

app = FastAPI(title="AutoVision - Gestão de Boletos API", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"], 
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class BoletoCreate(BaseModel):
    distribuidora: str = Field(..., min_length=1)
    numero_nota: Optional[str] = "S/N"
    valor_parcela: float = Field(..., gt=0)
    data_emissao: Optional[str] = None
    data_vencimento: str = Field(..., pattern=r"^\d{4}-\d{2}-\d{2}$")
    parcela_atual: int = Field(1, ge=1)
    total_parcelas: int = Field(1, ge=1)

class ParcelaRestaurar(BaseModel):
    id: int
    status_anterior: str

class ExcluirLotePayload(BaseModel):
    ids: List[int]

class RestaurarLotePayload(BaseModel):
    parcelas: List[ParcelaRestaurar]

@app.get("/status")
def home():
    return {"status": "AutoVision online!"}

def formatar_data_br(data_iso: str) -> str:
    if not data_iso or len(data_iso) != 10:
        return data_iso or ""
    p = data_iso.split('-')
    return f"{p[2]}/{p[1]}/{p[0]}"

@app.get("/boletos")
def obter_boletos(
    distribuidora: Optional[str] = Query(None),
    dias_vencimento: Optional[int] = Query(None),
    status: Optional[str] = Query("Aberto")
):
    boletos_raw = db.listar_boletos()
    boletos_ativos = [b for b in boletos_raw if b["status"] != 'Excluido']

    mapa_compras = {}
    for b in boletos_ativos:
        dist = (b["distribuidora"] or '').strip().upper()
        nota = (b["numero_nota"] or 'S/N').strip().upper()
        emissao = b["data_emissao"] or b["data_vencimento"]
        venc = b["data_vencimento"]
        valor = float(b["valor_parcela"])
        st = b["status"]
        chave = f"{dist}__{nota}__{b['data_emissao'] or 'SEM_EMISSAO'}"

        if chave not in mapa_compras:
            mapa_compras[chave] = {
                "inicio": emissao,
                "termino": venc,
                "total_parcelas": b["total_parcelas"],
                "valor_total": 0.0,
                "total_pago": 0.0,
                "total_pendente": 0.0,
                "qtd_pagas": 0,
                "todos_ids": [],
                "todas_parcelas_info": []
            }
        
        info = mapa_compras[chave]
        info["todos_ids"].append(b["id"])
        info["todas_parcelas_info"].append({"id": b["id"], "status_anterior": b["status"]})

        if venc > info["termino"]:
            info["termino"] = venc
        if emissao < info["inicio"]:
            info["inicio"] = emissao
            
        info["valor_total"] += valor
        if st == 'Pago':
            info["total_pago"] += valor
            info["qtd_pagas"] += 1
        else:
            info["total_pendente"] += valor

    boletos_formatados = []
    hoje = datetime.now().date()
    data_limite = hoje + timedelta(days=dias_vencimento) if dias_vencimento is not None else None

    for b in boletos_ativos:
        dist = (b["distribuidora"] or '').strip().upper()
        nota = (b["numero_nota"] or 'S/N').strip().upper()
        chave = f"{dist}__{nota}__{b['data_emissao'] or 'SEM_EMISSAO'}"
        dados_compra = mapa_compras.get(chave)

        item = {
            "id": b["id"],
            "distribuidora": b["distribuidora"],
            "numero_nota": b["numero_nota"] or 'S/N',
            "valor_parcela": b["valor_parcela"],
            "data_emissao": b["data_emissao"],
            "data_vencimento": b["data_vencimento"],
            "parcela_atual": b["parcela_atual"],
            "total_parcelas": b["total_parcelas"],
            "status": b["status"],
            "compra_chave": chave,
            "compra_resumo": {
                "data_inicio": dados_compra["inicio"] if dados_compra else b["data_vencimento"],
                "data_termino": dados_compra["termino"] if dados_compra else b["data_vencimento"],
                "total_parcelas": dados_compra["total_parcelas"] if dados_compra else b["total_parcelas"],
                "qtd_pagas": dados_compra["qtd_pagas"] if dados_compra else 0,
                "valor_total": round(dados_compra["valor_total"], 2) if dados_compra else round(b["valor_parcela"] * b["total_parcelas"], 2),
                "total_pago": round(dados_compra["total_pago"], 2) if dados_compra else 0.0,
                "total_pendente": round(dados_compra["total_pendente"], 2) if dados_compra else round(b["valor_parcela"], 2),
                "todos_ids_compra": dados_compra["todos_ids"] if dados_compra else [b["id"]],
                "todas_parcelas_info": dados_compra["todas_parcelas_info"] if dados_compra else [{"id": b["id"], "status_anterior": b["status"]}]
            }
        }
        
        if status and status.lower() != 'todos':
            if item["status"].lower() != status.lower():
                continue

        if distribuidora and distribuidora.strip():
            if distribuidora.strip().lower() not in item["distribuidora"].lower():
                continue

        if dias_vencimento is not None:
            try:
                data_venc = datetime.strptime(item["data_vencimento"], '%Y-%m-%d').date()
                if not (data_venc <= data_limite):
                    continue
            except ValueError:
                continue

        boletos_formatados.append(item)
        
    return boletos_formatados

@app.post("/boletos")
def criar_boleto(novo_boleto: BoletoCreate):
    if novo_boleto.parcela_atual > novo_boleto.total_parcelas:
        raise HTTPException(status_code=400, detail="A parcela inicial não pode ser superior ao total de parcelas.")

    db.cadastrar_boleto(
        novo_boleto.distribuidora, 
        novo_boleto.numero_nota, 
        novo_boleto.valor_parcela, 
        novo_boleto.data_emissao, 
        novo_boleto.data_vencimento, 
        novo_boleto.parcela_atual, 
        novo_boleto.total_parcelas
    )
    return {"mensagem": "Boleto e parcelas registradas com sucesso!"}

@app.put("/boletos/{boleto_id}/pagar")
def marcar_como_pago(boleto_id: int):
    with db.obter_conexao() as conexao:
        cursor = conexao.cursor()
        cursor.execute("UPDATE boletos SET status = 'Pago' WHERE id = ?", (boleto_id,))
        if cursor.rowcount == 0:
            raise HTTPException(status_code=404, detail="Boleto não encontrado.")
    return {"mensagem": f"Boleto ID {boleto_id} marcado como Pago!"}

@app.put("/boletos/{boleto_id}/reabrir")
def reabrir_boleto(boleto_id: int):
    """Permite desfazer a quitação de uma parcela marcada como paga por engano."""
    with db.obter_conexao() as conexao:
        cursor = conexao.cursor()
        cursor.execute("UPDATE boletos SET status = 'Aberto' WHERE id = ?", (boleto_id,))
        if cursor.rowcount == 0:
            raise HTTPException(status_code=404, detail="Boleto não encontrado.")
    return {"mensagem": f"Boleto ID {boleto_id} reaberto com sucesso!"}

@app.post("/compras/excluir")
def excluir_compra(dados: ExcluirLotePayload):
    if not dados.ids:
        raise HTTPException(status_code=400, detail="Nenhum boleto informado para exclusão.")
    
    placeholders = ','.join(['?'] * len(dados.ids))
    with db.obter_conexao() as conexao:
        cursor = conexao.cursor()
        cursor.execute(f"UPDATE boletos SET status = 'Excluido' WHERE id IN ({placeholders})", dados.ids)
    return {"mensagem": "Compra excluída com sucesso!"}

@app.post("/compras/restaurar")
def restaurar_compra(dados: RestaurarLotePayload):
    if not dados.parcelas:
        raise HTTPException(status_code=400, detail="Nenhuma parcela informada para restauração.")

    with db.obter_conexao() as conexao:
        cursor = conexao.cursor()
        for p in dados.parcelas:
            cursor.execute("UPDATE boletos SET status = ? WHERE id = ?", (p.status_anterior, p.id))
    return {"mensagem": "Compra restaurada com sucesso!"}

@app.get("/boletos/exportar/excel")
def exportar_excel(
    distribuidora: Optional[str] = Query(None),
    dias_vencimento: Optional[int] = Query(None),
    status: Optional[str] = Query("Aberto")
):
    boletos = obter_boletos(distribuidora=distribuidora, dias_vencimento=dias_vencimento, status=status)
    if not boletos:
        raise HTTPException(status_code=404, detail="Nenhum dado encontrado para exportar.")

    if TEM_OPENPYXL:
        wb = Workbook()
        ws = wb.active
        ws.title = "AutoVision Painel"
        
        # Correção segura para exibição de linhas de grade
        ws.sheet_view.showGridLines = True

        c_azul_profundo = "0B1329"
        c_azul_card = "0F172A"
        c_azul_cabecalho = "1E293B"
        c_ciano = "00D2FF"
        c_esmeralda = "10B981"
        c_branco = "F8FAFC"
        c_cinza_texto = "94A3B8"
        c_borda_card = "334155"

        fonte_titulo = Font(name="Segoe UI", size=15, bold=True, color=c_ciano)
        fonte_subtitulo = Font(name="Segoe UI", size=9, italic=True, color=c_cinza_texto)
        fonte_kpi_rotulo = Font(name="Segoe UI", size=9, bold=True, color=c_cinza_texto)
        fonte_kpi_valor_ciano = Font(name="Segoe UI", size=14, bold=True, color=c_ciano)
        fonte_kpi_valor_esmeralda = Font(name="Segoe UI", size=14, bold=True, color=c_esmeralda)
        fonte_kpi_valor_branco = Font(name="Segoe UI", size=14, bold=True, color=c_branco)
        
        fonte_cabecalho = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        fonte_dado = Font(name="Segoe UI", size=10)
        fonte_total = Font(name="Segoe UI", size=11, bold=True, color=c_azul_profundo)

        fill_titulo = PatternFill(start_color=c_azul_profundo, end_color=c_azul_profundo, fill_type="solid")
        fill_card = PatternFill(start_color=c_azul_card, end_color=c_azul_card, fill_type="solid")
        fill_cabecalho = PatternFill(start_color=c_azul_cabecalho, end_color=c_azul_cabecalho, fill_type="solid")
        fill_zebra = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        fill_pago = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        fill_aberto = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

        borda_fina = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        borda_card = Border(
            left=Side(style='thin', color=c_borda_card),
            right=Side(style='thin', color=c_borda_card),
            top=Side(style='thin', color=c_borda_card),
            bottom=Side(style='thin', color=c_borda_card)
        )
        borda_dupla_total = Border(
            top=Side(style='thin', color='000000'),
            bottom=Side(style='double', color='000000')
        )

        ws.merge_cells("A1:I1")
        cell_t = ws["A1"]
        cell_t.value = "⚡ AUTOVISION - PAINEL FINANCEIRO DE BOLETOS & COMPRAS"
        cell_t.font = fonte_titulo
        cell_t.fill = fill_titulo
        cell_t.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        ws.merge_cells("A2:I2")
        ws["A2"].value = f"Relatório Gerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M')} | Filtro Aplicado: Status {status or 'Todos'}"
        ws["A2"].font = fonte_subtitulo
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 18

        ws.row_dimensions[3].height = 6
        ws.row_dimensions[4].height = 16
        ws.row_dimensions[5].height = 24

        linha_inicio_dados = 8
        linha_fim_dados = linha_inicio_dados + len(boletos) - 1

        ws.merge_cells("A4:C4")
        ws["A4"].value = "💳 TOTAL EM ABERTO (PENDENTE)"
        ws["A4"].font = fonte_kpi_rotulo
        ws["A4"].alignment = Alignment(horizontal="center", vertical="center")
        
        ws.merge_cells("A5:C5")
        ws["A5"].value = f'=SUMIF(I{linha_inicio_dados}:I{linha_fim_dados}, "ABERTO", G{linha_inicio_dados}:G{linha_fim_dados})'
        ws["A5"].font = fonte_kpi_valor_ciano
        ws["A5"].number_format = '"R$" #,##0.00'
        ws["A5"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("D4:F4")
        ws["D4"].value = "🏁 TOTAL JÁ QUITADO"
        ws["D4"].font = fonte_kpi_rotulo
        ws["D4"].alignment = Alignment(horizontal="center", vertical="center")
        
        ws.merge_cells("D5:F5")
        ws["D5"].value = f'=SUMIF(I{linha_inicio_dados}:I{linha_fim_dados}, "PAGO", G{linha_inicio_dados}:G{linha_fim_dados})'
        ws["D5"].font = fonte_kpi_valor_esmeralda
        ws["D5"].number_format = '"R$" #,##0.00'
        ws["D5"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("G4:I4")
        ws["G4"].value = "📦 VOLUME DE PARCELAS LISTADAS"
        ws["G4"].font = fonte_kpi_rotulo
        ws["G4"].alignment = Alignment(horizontal="center", vertical="center")
        
        ws.merge_cells("G5:I5")
        ws["G5"].value = f'=COUNTA(A{linha_inicio_dados}:A{linha_fim_dados})'
        ws["G5"].font = fonte_kpi_valor_branco
        ws["G5"].alignment = Alignment(horizontal="center", vertical="center")

        for col_idx in range(1, 10):
            c4 = ws.cell(row=4, column=col_idx)
            c5 = ws.cell(row=5, column=col_idx)
            c4.fill = fill_card
            c5.fill = fill_card
            c4.border = borda_card
            c5.border = borda_card

        ws.row_dimensions[6].height = 10

        cabecalhos = [
            'Distribuidora', 'Nº Nota', 'Início Compra', 'Parcela', 
            'Vencimento', 'Término Compra', 'Valor Parcela', 'Total Compra', 'Status'
        ]
        
        ws.row_dimensions[7].height = 28
        for col_idx, cabecalho in enumerate(cabecalhos, start=1):
            cell = ws.cell(row=7, column=col_idx, value=cabecalho)
            cell.font = fonte_cabecalho
            cell.fill = fill_cabecalho
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = borda_fina

        ws.auto_filter.ref = f"A7:I{linha_fim_dados}"
        ws.freeze_panes = "A8"

        linha_atual = linha_inicio_dados
        for idx, b in enumerate(boletos):
            ws.row_dimensions[linha_atual].height = 21
            resumo = b['compra_resumo']
            
            ws.cell(row=linha_atual, column=1, value=b['distribuidora'])
            ws.cell(row=linha_atual, column=2, value=b['numero_nota'])
            ws.cell(row=linha_atual, column=3, value=formatar_data_br(resumo['data_inicio']))
            ws.cell(row=linha_atual, column=4, value=f"{b['parcela_atual']} de {b['total_parcelas']}")
            ws.cell(row=linha_atual, column=5, value=formatar_data_br(b['data_vencimento']))
            ws.cell(row=linha_atual, column=6, value=formatar_data_br(resumo['data_termino']))
            
            c_val = ws.cell(row=linha_atual, column=7, value=b['valor_parcela'])
            c_val.number_format = '"R$" #,##0.00'

            c_tot = ws.cell(row=linha_atual, column=8, value=resumo['valor_total'])
            c_tot.number_format = '"R$" #,##0.00'

            st_limpo = "PAGO" if b['status'] == 'Pago' else "ABERTO"
            c_status = ws.cell(row=linha_atual, column=9, value=st_limpo)

            for col_i in range(1, 10):
                celula = ws.cell(row=linha_atual, column=col_i)
                celula.font = fonte_dado
                celula.border = borda_fina
                if idx % 2 == 1:
                    celula.fill = fill_zebra

            ws.cell(row=linha_atual, column=1).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row=linha_atual, column=2).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=linha_atual, column=3).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=linha_atual, column=4).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=linha_atual, column=5).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=linha_atual, column=6).alignment = Alignment(horizontal="center", vertical="center")
            c_val.alignment = Alignment(horizontal="right", vertical="center")
            c_tot.alignment = Alignment(horizontal="right", vertical="center")
            c_status.alignment = Alignment(horizontal="center", vertical="center")

            if st_limpo == 'PAGO':
                c_status.fill = fill_pago
                c_status.font = Font(name="Segoe UI", size=10, bold=True, color="166534")
            else:
                c_status.fill = fill_aberto
                c_status.font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")

            linha_atual += 1

        ws.row_dimensions[linha_atual].height = 26
        ws.merge_cells(start_row=linha_atual, start_column=1, end_row=linha_atual, end_column=6)
        
        c_rotulo_total = ws.cell(row=linha_atual, column=1)
        c_rotulo_total.value = "TOTAL DAS PARCELAS (RECALCULA CONFORME OS FILTROS):"
        c_rotulo_total.font = fonte_total
        c_rotulo_total.alignment = Alignment(horizontal="right", vertical="center")

        c_subtotal_val = ws.cell(row=linha_atual, column=7)
        c_subtotal_val.value = f"=SUBTOTAL(109, G{linha_inicio_dados}:G{linha_fim_dados})"
        c_subtotal_val.font = Font(name="Segoe UI", size=11, bold=True, color="0070F3")
        c_subtotal_val.number_format = '"R$" #,##0.00'
        c_subtotal_val.alignment = Alignment(horizontal="right", vertical="center")

        c_vazio_h = ws.cell(row=linha_atual, column=8, value="-")
        c_vazio_h.alignment = Alignment(horizontal="center", vertical="center")

        for col_i in range(1, 10):
            ws.cell(row=linha_atual, column=col_i).border = borda_dupla_total

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in [1, 2, 4, 5, linha_atual]:
                    continue
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        headers = {
            'Content-Disposition': 'attachment; filename="AutoVision_Relatorio_Boletos.xlsx"',
            'Access-Control-Expose-Headers': 'Content-Disposition'
        }
        return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)

    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=';')
    writer.writerow(['Distribuidora', 'Nº Nota', 'Início Compra', 'Parcela', 'Vencimento', 'Término Compra', 'Valor Parcela (R$)', 'Total Compra (R$)', 'Status'])
    for b in boletos:
        resumo = b['compra_resumo']
        writer.writerow([
            b['distribuidora'], b['numero_nota'], formatar_data_br(resumo['data_inicio']),
            f"{b['parcela_atual']}/{b['total_parcelas']}", formatar_data_br(b['data_vencimento']),
            formatar_data_br(resumo['data_termino']),
            f"{b['valor_parcela']:.2f}".replace('.', ','),
            f"{resumo['valor_total']:.2f}".replace('.', ','),
            b['status']
        ])
    
    bytes_buffer = io.BytesIO(buffer.getvalue().encode('utf-8-sig'))
    headers = {
        'Content-Disposition': 'attachment; filename="AutoVision_Relatorio_Boletos.csv"',
        'Access-Control-Expose-Headers': 'Content-Disposition'
    }
    return StreamingResponse(bytes_buffer, media_type="text/csv", headers=headers)

caminho_dist = obter_diretorio_dist()
if caminho_dist.exists() and (caminho_dist / "index.html").exists():
    app.mount("/", StaticFiles(directory=str(caminho_dist), html=True), name="frontend")